from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = r"C:\Users\madir\Desktop\omoikiri-kz-new-final-beta\HOSTER-SUPPORT-2026-10-21.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")
        element.append(tag)
    tc_pr.append(element)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.68)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.88)
sec.right_margin = Inches(0.88)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.3)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

title = styles["Title"]
title.font.name = "Calibri"
title.font.size = Pt(17)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(7)
title_ppr = title._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)

for name in ("Heading 1", "Heading 2"):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(3)

doc.add_paragraph("Запуск Node.js на основном домене OMOIKIRI KZ", "Title")
meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(8)
meta.add_run("Заявка №2709962").bold = True
meta.add_run("  ·  21 октября 2026 г.  ·  omoikiri.kz")

doc.add_paragraph("Здравствуйте!")
doc.add_paragraph(
    "Проверили предложенную вами конфигурацию. Корневая папка сайта теперь находилась "
    "внутри папки Node.js-приложения, но запросы к основному домену всё равно не попадали в приложение. "
    "Просим проверить связку Apache, nginx и Passenger и указать рабочие параметры Plesk."
)

doc.add_paragraph("Что настроили", "Heading 1")
doc.add_paragraph(
    "Для omoikiri.kz установили корневую папку сайта /production-proxy/public, "
    "корневую папку Node.js-приложения /production-proxy и файл запуска app.js. "
    "Предупреждение Plesk о несовместимых путях исчезло. После нажатия «Включить Node.js» "
    "интерфейс сообщил, что приложение включено."
)

doc.add_paragraph("Результаты проверки", "Heading 1")
table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Inches(4.9)
table.columns[1].width = Inches(1.7)
for cell, value in zip(table.rows[0].cells, ("Адрес", "Результат")):
    cell.text = value
    shade(cell, "EAEAEA")
    for run in cell.paragraphs[0].runs:
        run.bold = True
for address, result in (
    ("https://omoikiri.kz/", "HTTP 403"),
    ("https://omoikiri.kz/healthz", "HTTP 404"),
    ("https://staging.omoikiri.kz/healthz", "HTTP 200"),
):
    row = table.add_row()
    row.cells[0].text = address
    row.cells[1].text = result
for row in table.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        borders(cell)
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            for run in para.runs:
                run.font.size = Pt(9.5)

doc.add_paragraph("Запись журнала Apache", "Heading 1")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.18)
p.paragraph_format.space_after = Pt(7)
r = p.add_run(
    "AH01276: Cannot serve directory "
    "/var/www/vhosts/omoikiri.kz/production-proxy/public/: "
    "No matching DirectoryIndex found, and server-generated directory index forbidden by Options directive"
)
r.font.name = "Consolas"
r.font.size = Pt(8.8)
doc.add_paragraph(
    "Судя по этой записи, Apache обрабатывал папку как статический сайт, "
    "а запросы не передавались Node.js/Passenger. Это наш вывод по журналу; "
    "просим подтвердить причину на стороне сервера."
)

doc.add_paragraph("Текущее состояние", "Heading 1")
doc.add_paragraph(
    "Сразу после проверки вернули корневую папку /httpdocs и отключили Node.js для основного домена. "
    "Временная переадресация HTTP 302 на работающий staging восстановлена и проверена. "
    "Файлы старого сайта и базу данных не меняли."
)

doc.add_paragraph("Что просим сделать", "Heading 1")
requests = [
    "Проверить сгенерированную конфигурацию Apache, nginx и Passenger для omoikiri.kz: "
    "почему запросы не доходят до app.js при корректном расположении document root.",
    "Проверить, поддерживается ли одновременная работа Node.js на omoikiri.kz и staging.omoikiri.kz "
    "в этой подписке и не сбрасываются ли настройки Node.js при сохранении параметров хостинга.",
    "Сообщить точные параметры Plesk или исправить конфигурацию сервера, чтобы основной домен "
    "обслуживал Node.js-приложение без смены DNS и без влияния на staging.",
    "Прислать точную запись упомянутой вами JavaScript-ошибки staging: время, путь или URL "
    "и полный текст ошибки, чтобы мы смогли её воспроизвести.",
]
for item in requests:
    doc.add_paragraph(item, "List Number")

doc.add_paragraph(
    "Пожалуйста, не отключайте текущую переадресацию основного домена и не переключайте "
    "публичный трафик без согласования. После вашего ответа мы повторим проверку главной, "
    "/healthz, карточки товара и оформления заказа на основном адресе."
)
doc.add_paragraph("Ответ через форму личного кабинета отправить не удалось: она возвращает HTTP 403.")
doc.add_paragraph("Спасибо!")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Заявка №2709962  ·  omoikiri.kz").font.size = Pt(8)

doc.save(OUT)
print(OUT)
